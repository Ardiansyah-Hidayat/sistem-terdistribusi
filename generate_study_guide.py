import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_bg(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=110, bottom=110, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_callout_box(doc, title, lines, bg_hex="F1F5F9"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(4)
    r_title = p0.add_run(title)
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for idx, line in enumerate(lines):
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3 if idx < len(lines) - 1 else 0)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(line)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    doc.add_paragraph()


def add_styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h_text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=130, right=130)

    for r_idx, row_data in enumerate(rows):
        row_cells = table.add_row().cells
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = ""
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(val))
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            set_cell_bg(row_cells[c_idx], bg)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=130, right=130)

    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    doc.add_paragraph()


def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.22
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.bold = True
        rb.font.name = "Calibri"
        rb.font.size = Pt(11)
        rb.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p


def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(bold_prefix)
    r1.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.name = "Calibri"
    r2.font.size = Pt(10.5)
    return p


def add_img(doc, path, caption, width_inches=5.7):
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width_inches))

    pc = doc.add_paragraph()
    pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pc.paragraph_format.space_before = Pt(0)
    pc.paragraph_format.space_after = Pt(8)
    rc = pc.add_run(caption)
    rc.italic = True
    rc.font.name = "Calibri"
    rc.font.size = Pt(9.5)
    rc.font.color.rgb = RGBColor(0x47, 0x55, 0x69)


def build_study_guide():
    doc = Document()

    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.95)
        section.right_margin = Inches(0.95)

    # ==================== HEADER ====================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    rt = p_title.add_run(
        "MODUL BELAJAR & PANDUAN PRESENTASI LENGKAP\n"
        "SISTEM PENYIMPANAN BERKAS TERDISTRIBUSI (MASTER-WORKER)"
    )
    rt.bold = True
    rt.font.name = "Calibri"
    rt.font.size = Pt(15)
    rt.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    rs = p_sub.add_run(
        "Pegangan Khusus Presenter: Konsep Inti, Naskah Bicara Per Gambar, Skenario Live Demo, "
        "Bedah Kode, & Bocoran Tanya-Jawab Dosen"
    )
    rs.italic = True
    rs.font.name = "Calibri"
    rs.font.size = Pt(10.5)
    rs.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # =========================================================================
    # BAGIAN 1: PAHAMI KONSEP INTI DALAM 3 MENIT
    # =========================================================================
    add_h1(doc, "BAGIAN 1 — PAHAMI KONSEP INTI DALAM 3 MENIT")

    add_para(
        doc,
        "Sebelum mulai presentasi, pastikan kamu memegang 4 ide utama ini. Seluruh sistem yang kita buat "
        "berputar di sekitar 4 poin ini:"
    )

    add_bullet(
        doc,
        "1. Masalah yang Diselesaikan (Single Point of Failure / SPOF): ",
        "Kalau kita menyimpan file hanya di 1 server (server tunggal), begitu server itu mati/rusak, maka "
        "seluruh file tidak bisa diakses. Sistem kita menyelesaikan masalah ini dengan menyebarkan salinan file "
        "ke beberapa node penyimpanan."
    )
    add_bullet(
        doc,
        "2. Arsitektur Master-Worker (1 Koordinator + 3 Storage Node): ",
        "Sistem kita memisahkan tugas 'pengatur' dan 'penyimpan'. Master Coordinator (port 8000) bertugas sebagai "
        "otak yang mengatur buku catatan lokasi file (metadata JSON), memeriksa kesehatan node, dan melayani pengguna. "
        "Sedangkan Storage Node 1, 2, dan 3 (port 8001, 8002, 8003) adalah pekerja yang menyimpan file fisik."
    )
    add_bullet(
        doc,
        "3. Replikasi Otomatis (Replication Factor = 2) dengan Algoritma Least-Loaded: ",
        "Setiap kali pengguna mengunggah 1 file, Koordinator otomatis memilih 2 Storage Node aktif yang paling "
        "sedikit isi filenya (Least-Loaded), lalu menyimpan salinan file tersebut ke kedua node itu. Jadi setiap "
        "file selalu punya 2 salinan di 2 node berbeda."
    )
    add_bullet(
        doc,
        "4. Heartbeat Monitoring & Automatic Failover: ",
        "Koordinator mengecek detak jantung (/health) ketiga node setiap 3 detik. Kalau Node 1 tiba-tiba mati saat "
        "pengguna ingin mengunduh file, pengguna TIDAK AKAN mengalami error karena Koordinator otomatis mengalihkan "
        "pengambilan file (Failover) ke salinan cadangan di Node 2 atau Node 3."
    )

    add_callout_box(
        doc,
        "ANALOGI SEDERHANA (Gunakan ini saat presentasi agar audiens langsung paham):",
        [
            "Bayangkan Master Coordinator seperti Resepsionis Perpustakaan yang memegang katalog buku, "
            "sedangkan Storage Node 1, 2, dan 3 adalah 3 Ruang Gudang Buku yang terpisah.",
            "• Saat kita menitipkan 1 buku baru (Upload), Resepsionis menggandakan buku tersebut menjadi 2 salinan "
            "dan menyimpannya di 2 gudang yang paling luang (misal Gudang 1 dan Gudang 2).",
            "• Saat kita ingin meminjam buku itu kembali (Download), Resepsionis mengecek catatan: 'Buku ini ada di "
            "Gudang 1 dan Gudang 2'. Kalau pintu Gudang 1 ternyata terkunci/rusak (Node Mati), Resepsionis langsung "
            "mengambilkan buku yang sama dari Gudang 2 tanpa membuat peminjam gagal mendapatkan bukunya."
        ]
    )

    # =========================================================================
    # BAGIAN 2: NASKAH BICARA (SPEAKER SCRIPT) PER DIAGRAM & SCREENSHOT UI
    # =========================================================================
    add_h1(doc, "BAGIAN 2 — NASKAH BICARA (SPEAKER SCRIPT) PER SLIDE / GAMBAR")
    add_para(
        doc,
        "Bagian ini berisi panduan kalimat yang bisa langsung kamu ucapkan saat menampilkan masing-masing "
        "diagram dan screenshot UI di slide PPT."
    )

    # --- Slide 1: Diagram Arsitektur ---
    add_h2(doc, "2.1 Saat Menampilkan Diagram Arsitektur Sistem (Master-Worker)")
    add_img(doc, os.path.join("diagrams", "arsitektur_sistem.png"), "Diagram 1: Arsitektur Sistem Master-Worker", width_inches=5.6)
    add_callout_box(
        doc,
        "Kalimat yang Bisa Kamu Ucapkan:",
        [
            "\"Pada gambar arsitektur ini, sistem kami menggunakan pola Master-Worker. Di sebelah kiri adalah Klien "
            "(Web Browser) yang hanya berkomunikasi ke satu pintu, yaitu Master Coordinator di port 8000. Klien tidak "
            "perlu tahu di node mana file disimpan (ini disebut Location Transparency).\"",
            "\"Di bagian tengah, Master Coordinator menjalankan 5 tugas utama: mencatat katalog metadata di file JSON, "
            "memilih node dengan beban terendah (Least-Loaded Selector), memantau kesehatan node lewat Heartbeat setiap "
            "3 detik, mengatur hak akses Admin dan Tamu (RBAC), serta melakukan Failover saat ada node yang mati.\"",
            "\"Di sebelah kanan terdapat 3 Storage Node independen (Node 1 di port 8001, Node 2 di port 8002, dan "
            "Node 3 di port 8003) yang bertugas menyimpan dan mengirimkan file fisik saat diperintahkan oleh Koordinator.\""
        ]
    )

    # --- Slide 2: Diagram Alur Replikasi & Failover ---
    add_h2(doc, "2.2 Saat Menampilkan Diagram Alur Replikasi dan Failover Otomatis")
    add_img(doc, os.path.join("diagrams", "alur_failover_replikasi.png"), "Diagram 2: Alur Replikasi (A) dan Failover Otomatis (B)", width_inches=5.8)
    add_callout_box(
        doc,
        "Kalimat yang Bisa Kamu Ucapkan:",
        [
            "\"Diagram ini menjelaskan dua mekanisme utama sistem kami. Pertama pada Panel A (sebelah kiri) adalah "
            "alur saat pengguna mengunggah file. File dikirim ke Koordinator melalui POST /api/upload. Koordinator "
            "kemudian mengecek node mana yang aktif dan memiliki jumlah file paling sedikit, lalu memilih 2 node "
            "(karena Replication Factor kita bernilai 2). File kemudian dikirim ke Storage Node 1 dan Storage Node 2, "
            "lalu lokasinya disimpan di katalog metadata JSON.\"",
            "\"Kedua, pada Panel B (sebelah kanan) adalah alur Automatic Failover saat mengunduh file. Ketika klien "
            "meminta unduhan lewat GET /api/download/{id}, Koordinator melihat bahwa file tersebut ada di node1 dan "
            "node2. Apabila Storage Node 1 sedang mati atau gagal dihubungi (ditunjukkan garis putus-putus), Koordinator "
            "otomatis mengalihkan pengambilan file ke Storage Node 2 yang masih aktif, lalu mengirimkan file tersebut "
            "ke pengguna beserta informasi header X-Served-By: node2.\""
        ]
    )

    # --- Slide 3: Screenshot Halaman Login ---
    add_h2(doc, "2.3 Saat Menampilkan Screenshot 1: Halaman Login")
    add_img(doc, os.path.join("screenshots", "1_halaman_login.png"), "Screenshot 1: Halaman Login Konsol Web", width_inches=5.4)
    add_callout_box(
        doc,
        "Kalimat yang Bisa Kamu Ucapkan:",
        [
            "\"Ini adalah tampilan awal saat pengguna membuka aplikasi web kami. Kami menerapkan Role-Based Access "
            "Control (RBAC) dengan dua peran: Administrator dan Tamu (Guest).\"",
            "\"Administrator masuk menggunakan username dan password untuk mendapatkan hak kontrol penuh, termasuk "
            "mematikan/menyalakan node dan mengatur izin tamu. Sedangkan di sebelah kanan, pengguna Tamu cukup "
            "mengklik satu tombol 'Masuk sebagai tamu' jika akses tamu sedang dibuka oleh Administrator.\""
        ]
    )

    # --- Slide 4: Screenshot Dashboard Admin Normal ---
    add_h2(doc, "2.4 Saat Menampilkan Screenshot 2: Dashboard Administrator (Kondisi Normal)")
    add_img(doc, os.path.join("screenshots", "2_dashboard_admin_normal.png"), "Screenshot 2: Dashboard Administrator Kondisi Normal", width_inches=5.6)
    add_callout_box(
        doc,
        "Kalimat yang Bisa Kamu Ucapkan:",
        [
            "\"Setelah masuk sebagai Administrator, kita melihat dashboard operasional dalam kondisi normal. Di bagian "
            "paling atas ada panel 'Kontrol akses' khusus Admin untuk mengizinkan/memblokir akses tamu secara real-time "
            "dan mengaktifkan mode eksperimen tamu.\"",
            "\"Di bawahnya terdapat ringkasan kondisi klaster: Node aktif 3 / 3, Faktor replikasi 2, Jumlah berkas 3, "
            "dan Berkas tersedia 3 / 3.\"",
            "\"Perhatikan tabel Node Penyimpanan dan tabel Berkas di bawah: karena kita mengunggah 3 berkas dengan "
            "Faktor Replikasi 2 (total 6 replika), algoritma Least-Loaded kita membagi beban secara sangat rata: "
            "Node 1 menyimpan 2 replika, Node 2 menyimpan 2 replika, dan Node 3 menyimpan 2 replika. Di kolom "
            "'Lokasi replika' pada tabel bawah, kita bisa melihat pasangan node tempat tiap file disimpan: berkas "
            "pertama di (node1, node2), berkas kedua di (node3, node1), dan berkas ketiga di (node2, node3).\""
        ]
    )

    # --- Slide 5: Screenshot Simulasi Node Mati & Failover ---
    add_h2(doc, "2.5 Saat Menampilkan Screenshot 3: Simulasi Kegagalan Node & Failover")
    add_img(doc, os.path.join("screenshots", "3_simulasi_node_mati_failover.png"), "Screenshot 3: Simulasi Node 1 Mati dan Unduhan Failover dari Node 2", width_inches=5.6)
    add_callout_box(
        doc,
        "Kalimat yang Bisa Kamu Ucapkan:",
        [
            "\"Selanjutnya kami menguji ketahanan sistem terhadap kegagalan node (Fault Tolerance). Kami menekan "
            "tombol 'Matikan' pada Node 1 untuk menyimulasikan server pertama mengalami crash atau mati.\"",
            "\"Terlihat perubahan pada antarmuka: pertama, indikator Node Aktif turun menjadi 2 / 3 dan status "
            "Node 1 berubah menjadi merah ('Mati'). Kedua, pada tabel berkas di bawah, tulisan 'node1' otomatis "
            "dicoret untuk menandakan bahwa salinan di node1 sedang offline.\"",
            "\"Namun perhatikan kartu 'Berkas tersedia' di kanan atas: nilainya tetap 3 / 3 (100% file tetap aman)! "
            "Ketika kami mengklik tombol 'Unduh' pada file Laporan_Sistem_Terdistribusi.pdf yang awalnya ada di "
            "node1 dan node2, sistem otomatis melakukan Failover dan mengambil file tersebut dari node2 seperti "
            "terlihat pada notifikasi hijau di pojok kanan bawah.\""
        ]
    )

    # --- Slide 6: Screenshot Mode Tamu ---
    add_h2(doc, "2.6 Saat Menampilkan Screenshot 4: Tampilan Mode Tamu (Guest Mode)")
    add_img(doc, os.path.join("screenshots", "4_tampilan_mode_tamu.png"), "Screenshot 4: Tampilan Mode Tamu (Hanya Lihat & Unggah)", width_inches=5.4)
    add_callout_box(
        doc,
        "Kalimat yang Bisa Kamu Ucapkan:",
        [
            "\"Terakhir, ini adalah tampilan ketika pengguna masuk sebagai Tamu (Guest) pada mode standar. Panel "
            "Kontrol Akses, tombol 'Matikan' node, dan tombol 'Hapus' berkas disembunyikan seluruhnya agar tamu "
            "tidak bisa merusak klaster. Tamu hanya diizinkan melihat status node, mengunggah berkas baru, dan "
            "mengunduh berkas yang ada. Namun jika Admin menyalakan sakelar 'Mode eksperimen tamu', barulah tamu "
            "diberi izin sementara untuk ikut mencoba simulasi mematikan node.\""
        ]
    )

    # =========================================================================
    # BAGIAN 3: URUTAN SKENARIO LIVE DEMO (STEP-BY-STEP)
    # =========================================================================
    add_h1(doc, "BAGIAN 3 — PANDUAN URUTAN LIVE DEMO SAAT PRESENTASI")
    add_para(
        doc,
        "Jika dosen meminta kamu mendemokan aplikasinya secara langsung, ikuti 5 langkah berurutan ini "
        "(hanya butuh waktu sekitar 2–3 menit):"
    )

    add_styled_table(
        doc,
        headers=["Langkah", "Aksi yang Kamu Lakukan", "Apa yang Ditunjukkan ke Dosen"],
        rows=[
            [
                "1. Nyalakan Klaster",
                "Klik 2x file START_LOCAL.bat lalu buka browser ke http://localhost:8000",
                "Tunjukkan bahwa 1 Koordinator (8000) dan 3 Storage Node (8001–8003) langsung berjalan."
            ],
            [
                "2. Login Admin",
                "Masukkan username: ardi dan password: ardi123 lalu klik Masuk",
                "Tunjukkan status awal klaster: Node aktif 3 / 3 dan Faktor Replikasi = 2."
            ],
            [
                "3. Demo Replikasi (Upload)",
                "Seret/pilih 2 atau 3 file bebas ke kotak unggah",
                "Tunjukkan kolom 'Lokasi replika' di tabel bawah: setiap file otomatis masuk ke 2 node berbeda secara merata."
            ],
            [
                "4. Demo Node Mati & Failover",
                "Klik tombol 'Matikan' pada Node 1, lalu klik 'Unduh' pada salah satu file yang punya replika di node1",
                "Tunjukkan node1 tercoret, status Berkas Tersedia tetap penuh, dan muncul toast hijau 'diunduh dari node2/node3'."
            ],
            [
                "5. Demo Pemulihan & RBAC",
                "Klik 'Nyalakan' kembali pada Node 1, lalu coba matikan sakelar 'Izinkan akses tamu'",
                "Tunjukkan Node 1 kembali Aktif (3/3) dan jelaskan fitur kontrol akses real-time untuk memutus sesi tamu."
            ],
        ],
        col_widths=[1.2, 2.4, 2.8]
    )

    # =========================================================================
    # BAGIAN 4: BEDAH KODE INTI (JIKA DOSEN MINTA BUKA SOURCE CODE)
    # =========================================================================
    add_h1(doc, "BAGIAN 4 — PETA SOURCE CODE (JIKA DOSEN MINTA BUKA KODE)")
    add_para(
        doc,
        "Kalau dosen bertanya: 'Coba tunjukkan di bagian kode mana logika replikasi atau failover-nya?', "
        "kamu cukup membuka salah satu dari 4 file utama berikut:"
    )

    add_styled_table(
        doc,
        headers=["Nama File", "Peran File", "Fungsi / Bagian Penting yang Ditunjukkan"],
        rows=[
            [
                "config.py",
                "Konfigurasi Klaster & Topologi",
                "• CLUSTER_MODE ('local' vs 'vm')\n• STORAGE_NODES (daftar IP & port node1, node2, node3)\n• DEFAULT_REPLICATION_FACTOR = 2\n• HEARTBEAT_INTERVAL = 3 detik"
            ],
            [
                "coordinator/node_tracker.py",
                "Pemantau Kesehatan & Seleksi Node",
                "• check_node_health(): mengecek GET /health tiap node\n• select_replica_nodes(count=2): mengurutkan node online berdasarkan file_count paling sedikit (Least-Loaded)\n• toggle_node_process(): mematikan/menyalakan node untuk simulasi"
            ],
            [
                "coordinator/main.py",
                "Otak Koordinator (REST API Utama)",
                "• upload_file() [POST /api/upload]: memilih 2 node lalu mengirim file ke POST /store/{file_id}\n• download_file() [GET /api/download/{file_id}]: melakukan perulangan (loop) ke node replika yang online; jika node pertama gagal, otomatis pindah ke node kedua (Failover)"
            ],
            [
                "node/main.py",
                "Program Worker Storage Node",
                "• health_check() [GET /health]: menghitung jumlah file & ukuran byte di folder node\n• store_file() [POST /store/{file_id}]: menyimpan file fisik ke disk\n• retrieve_file() [GET /retrieve/{file_id}]: mengirim file fisik kembali ke Koordinator"
            ],
        ],
        col_widths=[1.6, 1.6, 3.2]
    )

    # =========================================================================
    # BAGIAN 5: BOCORAN PERTANYAAN DOSEN & JAWABAN JITUNYA (Q&A CHEAT SHEET)
    # =========================================================================
    add_h1(doc, "BAGIAN 5 — DAFTAR PERTANYAAN DOSEN & JAWABAN JITUNYA (Q&A)")

    qa_list = [
        (
            "Q1: Mengapa kalian memilih Faktor Replikasi (Replication Factor) = 2, bukan 1 atau 3?",
            "Jawaban: Kalau RF = 1, artinya tidak ada cadangan (sama saja tidak tahan kegagalan bila node tersebut mati). "
            "Kalau RF = 3 pada klaster berisi 3 node, setiap file akan disimpan di semua node sehingga memakan ruang disk "
            "100% di setiap node. Dengan RF = 2 pada 3 node, kita mendapatkan titik keseimbangan (trade-off) terbaik antara "
            "efisiensi kapasitas penyimpanan dan ketersediaan tinggi (High Availability): sistem tetap tahan apabila 1 node "
            "mati total namun hemat ruang penyimpanan sebesar 33% dibanding menyalin ke seluruh node."
        ),
        (
            "Q2: Bagaimana algoritma Koordinator dalam memilih 2 node saat ada file yang diunggah?",
            "Jawaban: Kami menggunakan algoritma Least-Loaded Node Selection pada fungsi select_replica_nodes(). "
            "Pertama, Koordinator menyaring hanya node yang berstatus 'online'. Kedua, node yang online diurutkan "
            "berdasarkan jumlah berkas (file_count) dari yang paling sedikit ke paling banyak. Apabila jumlah berkasnya "
            "sama, urutannya diacak terlebih dahulu lalu dipilih 2 teratas. Dengan begitu, distribusi berkas di ketiga "
            "node selalu seimbang."
        ),
        (
            "Q3: Bagaimana jika sebuah node baru saja mati 1 detik yang lalu (sebelum jadwal Heartbeat 3 detik berikutnya), lalu ada pengguna yang mengunduh file?",
            "Jawaban: Sistem kami tetap aman karena menerapkan berlapis dua mekanisme: Proactive Heartbeat (setiap 3 detik) "
            "dan Reactive/Lazy Failover saat unduhan berlangsung. Di dalam fungsi download_file(), Koordinator menggunakan "
            "blok try-except di dalam perulangan (for node_id in available_replicas). Jadi meskipun node pertama masih "
            "tercatat online namun ternyata gagal dihubungi saat detik itu juga, blok except akan menangkap error koneksi "
            "tersebut dan langsung lanjut mengambil file dari node kandidat berikutnya."
        ),
        (
            "Q4: Apa yang terjadi jika 2 Storage Node dimatikan sekaligus (misalnya Node 1 dan Node 2 mati, hanya Node 3 yang hidup)?",
            "Jawaban: Karena setiap file memiliki 2 replika, maka:\n"
            "• File yang salah satu replikanya berada di Node 3 (misalnya disimpan di [node1, node3] atau [node2, node3]) "
            "TETAP BISA DIUNDUH 100% dari Node 3.\n"
            "• Sedangkan file yang kebetulan kedua replikanya berada di [node1, node2] akan ditandai pada tabel sebagai "
            "'Semua replika sedang mati' dan tombol unduhnya dinonaktifkan sementara (HTTP 503) sampai salah satu dari "
            "Node 1 atau Node 2 dinyalakan kembali."
        ),
        (
            "Q5: Ketika Node 1 yang tadinya mati dinyalakan kembali, apakah file di dalamnya hilang?",
            "Jawaban: Tidak hilang. Karena setiap Storage Node menyimpan berkas secara persisten di dalam direktori disk "
            "masing-masing (storage_data/node1, node2, node3) dan katalog metadata disimpan di coordinator_metadata.json, "
            "maka begitu Node 1 dinyalakan kembali, dalam waktu maksimal 3 detik Heartbeat akan mendeteksi Node 1 kembali "
            "'Aktif' beserta seluruh file di dalamnya."
        ),
        (
            "Q6: Apa kelemahan dari arsitektur Master-Worker yang kalian buat saat ini?",
            "Jawaban: Sesuai teori sistem terdistribusi, kelemahan klasik arsitektur Master-Worker dengan satu Master "
            "adalah jika Master Coordinator-nya sendiri yang mati, maka klien tidak bisa mengakses katalog metadata "
            "meskipun ketiga Storage Node masih hidup. Untuk pengembangan selanjutnya, solusinya adalah menambahkan "
            "Secondary/Standby Coordinator dengan sinkronisasi metadata menggunakan protokol konsensus seperti Raft "
            "atau Apache ZooKeeper (seperti pada arsitektur HDFS High Availability)."
        ),
    ]

    for q_text, a_text in qa_list:
        add_callout_box(doc, q_text, [a_text], bg_hex="F8FAFC")

    out_name = "Panduan_Belajar_dan_Presentasi_Sistem_Terdistribusi.docx"
    doc.save(out_name)
    print(f"[OK] Saved {out_name}")


if __name__ == "__main__":
    build_study_guide()
