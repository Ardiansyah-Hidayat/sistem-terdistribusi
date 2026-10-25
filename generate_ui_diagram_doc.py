import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_section_item(doc, subheading, image_path, caption, explanation_paragraphs, width_inches=6.0):
    # 1. Sub-heading
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(14)
    p_sub.paragraph_format.space_after = Pt(6)
    p_sub.paragraph_format.keep_with_next = True
    r_sub = p_sub.add_run(subheading)
    r_sub.bold = True
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(12.5)
    r_sub.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # 2. Gambar
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(8)
        r_cap = p_cap.add_run(caption)
        r_cap.italic = True
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # 3. Penjelasan Singkat
    for idx, text in enumerate(explanation_paragraphs):
        p_exp = doc.add_paragraph()
        p_exp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_exp.paragraph_format.line_spacing = 1.25
        p_exp.paragraph_format.space_before = Pt(0)
        p_exp.paragraph_format.space_after = Pt(6 if idx < len(explanation_paragraphs) - 1 else 12)
        if idx == 0:
            r_label = p_exp.add_run("Penjelasan Singkat: ")
            r_label.bold = True
            r_label.font.name = "Times New Roman"
            r_label.font.size = Pt(11)
        r_txt = p_exp.add_run(text)
        r_txt.font.name = "Times New Roman"
        r_txt.font.size = Pt(11)
        r_txt.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def build_ui_and_diagram_doc():
    doc = Document()

    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(
        "DOKUMENTASI ANTARMUKA PENGGUNA (UI) DAN DIAGRAM SISTEM\n"
        "KLASTER PENYIMPANAN BERKAS TERDISTRIBUSI (MASTER-WORKER)"
    )
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(13.5)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(
        "Dokumen Pendamping Presentasi (PPT) dan Artikel Ilmiah"
    )
    r_sub.italic = True
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(10.5)

    # -------------------------------------------------------------------------
    # 1. Halaman Login
    # -------------------------------------------------------------------------
    add_section_item(
        doc,
        subheading="1. Halaman Login",
        image_path=os.path.join("screenshots", "1_halaman_login.png"),
        caption="Gambar 1. Antarmuka Halaman Login (Administrator dan Akses Tamu).",
        explanation_paragraphs=[
            "Halaman Login merupakan gerbang autentikasi berbasis peran (Role-Based Access Control) sebelum pengguna "
            "mengakses konsol klaster penyimpanan. Halaman ini terbagi menjadi dua bagian utama: formulir masuk "
            "Administrator menggunakan nama pengguna dan kata sandi untuk memperoleh hak kontrol penuh atas klaster, "
            "serta tombol 'Masuk sebagai tamu' yang memungkinkan akses cepat satu klik apabila fitur akses tamu "
            "sedang diizinkan oleh administrator."
        ],
        width_inches=5.8
    )

    # -------------------------------------------------------------------------
    # 2. Dashboard Administrator (Kondisi Normal)
    # -------------------------------------------------------------------------
    add_section_item(
        doc,
        subheading="2. Dashboard Administrator (Kondisi Normal)",
        image_path=os.path.join("screenshots", "2_dashboard_admin_normal.png"),
        caption="Gambar 2. Tampilan Dashboard Administrator Saat Seluruh Storage Node Aktif.",
        explanation_paragraphs=[
            "Dashboard Administrator menampilkan pusat kendali operasional klaster ketika ketiga simpul penyimpanan "
            "(Node 1, Node 2, dan Node 3) berada dalam kondisi aktif (Node aktif: 3 / 3). Pada bagian teratas terdapat "
            "panel 'Kontrol akses' khusus administrator yang berisi dua sakelar waktu nyata: 'Izinkan akses tamu' dan "
            "'Mode eksperimen tamu'. Di bawahnya terdapat ringkasan metrik klaster, tabel 'Node penyimpanan' yang "
            "memperlihatkan beban replika secara merata (masing-masing 2 replika per node hasil algoritma Least-Loaded "
            "dengan Faktor Replikasi = 2), area unggah berkas (drag-and-drop), serta tabel daftar berkas beserta "
            "lokasi simpul replikanya."
        ],
        width_inches=5.9
    )

    # -------------------------------------------------------------------------
    # 3. Simulasi Kegagalan Node dan Failover Otomatis
    # -------------------------------------------------------------------------
    add_section_item(
        doc,
        subheading="3. Simulasi Kegagalan Node dan Failover Otomatis",
        image_path=os.path.join("screenshots", "3_simulasi_node_mati_failover.png"),
        caption="Gambar 3. Simulasi Saat Node 1 Dimatikan dan Pengalihan Unduhan (Failover) ke Node 2.",
        explanation_paragraphs=[
            "Tampilan ini memperlihatkan mekanisme toleransi kegagalan (fault tolerance) ketika Node 1 dimatikan melalui "
            "tombol simulasi. Indikator 'Node aktif' turun menjadi 2 / 3, baris Node 1 berubah status menjadi 'Mati', "
            "dan label 'node1' pada kolom Lokasi replika dicoret secara otomatis untuk menandakan bahwa salinan pada "
            "node tersebut sedang tidak dapat dijangkau. Meskipun Node 1 mati, metrik 'Berkas tersedia' tetap "
            "menunjukkan 3 / 3 (100% berkas aman) karena setiap berkas masih memiliki salinan kedua di Node 2 atau "
            "Node 3. Ketika berkas diunduh, koordinator secara otomatis mengalihkan pengambilan data (automatic failover) "
            "ke node replika cadangan yang masih aktif sebagaimana terlihat pada notifikasi sudut kanan bawah."
        ],
        width_inches=5.9
    )

    # -------------------------------------------------------------------------
    # 4. Tampilan Mode Tamu (Guest View)
    # -------------------------------------------------------------------------
    add_section_item(
        doc,
        subheading="4. Tampilan Mode Tamu (Guest Mode)",
        image_path=os.path.join("screenshots", "4_tampilan_mode_tamu.png"),
        caption="Gambar 4. Tampilan Konsol pada Mode Tamu (Hanya Lihat dan Unggah).",
        explanation_paragraphs=[
            "Pada Mode Tamu standar (ketika 'Mode eksperimen tamu' tidak diaktifkan oleh administrator), antarmuka "
            "menyembunyikan panel Kontrol Akses, tombol simulasi 'Matikan/Nyalakan' node, serta tombol 'Hapus' berkas. "
            "Pengguna tamu tetap dapat memantau status kesehatan ketiga Storage Node secara langsung, mengunggah berkas "
            "baru ke dalam klaster, serta mengunduh berkas yang tersedia tanpa risiko mengubah konfigurasi atau "
            "mengganggu ketersediaan klaster."
        ],
        width_inches=5.8
    )

    # -------------------------------------------------------------------------
    # 5. Diagram Arsitektur Sistem (Master-Worker)
    # -------------------------------------------------------------------------
    add_section_item(
        doc,
        subheading="5. Diagram Arsitektur Sistem (Master-Worker)",
        image_path=os.path.join("diagrams", "arsitektur_sistem.png"),
        caption="Gambar 5. Diagram Arsitektur Sistem Penyimpanan Terdistribusi Berbasis Master-Worker.",
        explanation_paragraphs=[
            "Diagram arsitektur ini menggambarkan pemisahan tanggung jawab antara simpul pengendali (Master Coordinator "
            "pada port 8000) dan tiga simpul pekerja (Storage Node 1, 2, dan 3 pada port 8001–8003). Klien hanya perlu "
            "berkomunikasi dengan Master Coordinator melalui protokol HTTP REST API. Koordinator bertugas mengelola "
            "katalog metadata berkas (JSON), menjalankan pemeriksaan kesehatan simpul (heartbeat polling setiap 3 detik), "
            "memilih simpul penyimpanan dengan beban terendah (least-loaded selector), serta mengatur rute pengalihan "
            "unduhan (failover router) ke simpul pekerja baik pada mode klaster lokal (127.0.0.1) maupun mode mesin "
            "virtual VirtualBox (192.168.56.102–104)."
        ],
        width_inches=6.0
    )

    # -------------------------------------------------------------------------
    # 6. Diagram Alur Replikasi dan Failover Otomatis
    # -------------------------------------------------------------------------
    add_section_item(
        doc,
        subheading="6. Diagram Alur Replikasi dan Failover Otomatis",
        image_path=os.path.join("diagrams", "alur_failover_replikasi.png"),
        caption="Gambar 6. Diagram Alur Replikasi Berkas (Panel A) dan Failover Unduhan Otomatis (Panel B).",
        explanation_paragraphs=[
            "Diagram ini membandingkan dua alur kerja utama sistem. Pada Panel A (Alur Replikasi Berkas), setiap berkas "
            "yang diunggah klien ke koordinator akan diteruskan secara paralel ke 2 Storage Node aktif yang memiliki "
            "jumlah simpanan berkas paling sedikit (Replication Factor = 2), kemudian lokasi kedua replika dicatat ke "
            "dalam katalog metadata. Pada Panel B (Alur Failover Otomatis), ketika klien mengunduh berkas namun simpul "
            "penyimpan utama (Storage Node 1) dalam kondisi mati atau gagal dihubungi, koordinator secara otomatis "
            "membatalkan koneksi ke Node 1 dan mengambil berkas dari simpul replika cadangan (Storage Node 2) sehingga "
            "berkas tetap berhasil dikirim ke klien tanpa kegagalan layanan."
        ],
        width_inches=6.1
    )

    out_path = "Dokumentasi_UI_dan_Diagram_Sistem.docx"
    doc.save(out_path)
    print(f"[OK] Saved {out_path}")


if __name__ == "__main__":
    build_ui_and_diagram_doc()
